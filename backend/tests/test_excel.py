"""Excel 模板与导入解析测试。"""
import io

from openpyxl import Workbook, load_workbook

from app.services.excel_import import build_template, parse_excel


def _rows_of(content: bytes, sheet: str):
    ws = load_workbook(io.BytesIO(content), read_only=True)[sheet]
    return list(ws.iter_rows(values_only=True))


def test_template_downloadable_and_fills_back():
    content = build_template()
    rows = _rows_of(content, "数据流台账")
    assert rows[0][0] == "源系统"
    assert rows[1][0] == "订单中心"  # 示例行


def test_parse_creates_system_and_flow():
    wb = Workbook()
    ws = wb.active
    ws.title = "数据流台账"
    ws.append(["源系统", "源系统类型(系统/外部方)", "目标系统", "目标系统类型(系统/外部方)",
               "流转方式(接口/库表同步/文件/人工)", "数据内容(库.表 或 说明)",
               "字段映射(源字段->目标字段,逗号分隔)", "频率(实时/小时/日/月)", "负责人", "备注"])
    ws.append(["订单中心", "系统", "顺丰物流", "外部方", "接口", "发货单", "", "实时", "张三", ""])
    ws.append(["订单库", "系统", "数据仓库", "系统", "库表同步", "orders.orders",
               "order_id->id,amount->amount", "日", "李四", ""])
    buf = io.BytesIO()
    wb.save(buf)

    nodes, edges, errors = parse_excel(buf.getvalue())
    assert errors == []
    node_ids = {n.id for n in nodes}
    assert "sys:订单中心" in node_ids
    assert "sys:顺丰物流" in node_ids
    assert next(n for n in nodes if n.id == "sys:顺丰物流").type == "external"
    assert "tbl:orders.orders" in node_ids  # 数据内容是 库.表 → 建表节点

    flow = next(e for e in edges if e.id.startswith("flow:sys:订单中心->sys:顺丰物流"))
    assert flow.attrs["frequency"] == "实时"
    assert flow.attrs["owner"] == "张三"
    produce = next(e for e in edges if e.id.startswith("produces:sys:订单库->tbl:orders.orders"))
    assert produce.target_id == "tbl:orders.orders"
    consume = next(e for e in edges if e.id.startswith("consumes:tbl:orders.orders->"))
    assert consume.source_id == "tbl:orders.orders"


def test_parse_invalid_rows_reported():
    wb = Workbook()
    ws = wb.active
    ws.title = "数据流台账"
    ws.append(["源系统", "源系统类型(系统/外部方)", "目标系统", "目标系统类型(系统/外部方)",
               "流转方式(接口/库表同步/文件/人工)", "数据内容(库.表 或 说明)",
               "字段映射(源字段->目标字段,逗号分隔)", "频率(实时/小时/日/月)", "负责人", "备注"])
    ws.append(["只有源系统", "系统", "", "", "", "", "", "", "", ""])
    ws.append(["", "", "", "", "", "", "", "", "", ""])  # 空行忽略
    buf = io.BytesIO()
    wb.save(buf)

    nodes, edges, errors = parse_excel(buf.getvalue())
    assert len(errors) == 1
    assert nodes == [] and edges == []


def test_parse_wrong_sheet_rejected():
    wb = Workbook()
    wb.active.title = "错误的名字"
    buf = io.BytesIO()
    wb.save(buf)
    try:
        parse_excel(buf.getvalue())
        raise AssertionError("应当抛出 ValueError")
    except ValueError as e:
        assert "模板" in str(e)
