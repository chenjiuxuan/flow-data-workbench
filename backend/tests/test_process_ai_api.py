"""POST /api/process/ai/draft：mock 模型，零写库，错误码与现有保存接口。"""
import asyncio
import io
import json
import time

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api import router as data_router
from app.db import Base, get_session
from app.process_ai.errors import NOT_A_PROCESS, PARSE_FAILED, SCANNED_PDF
from app.process_ai.schema import GenDraft
from app.process_schemas import ProcessCreate
from app.process_api import router as process_router
from app.process_ai.api import router as ai_router
from tests.test_process_ai_extract import _blank_pdf, _pdf_with_text

SOURCE = "张三提交申请后，判断是否通过。通过则由李四执行办理，驳回则结束。并核对订单中心。"


@pytest.fixture(autouse=True)
def clear_llm_env(monkeypatch):
    for key in ("XILIU_LLM_BASE_URL", "XILIU_LLM_API_KEY", "XILIU_LLM_MODEL", "XILIU_LLM_FIXTURE"):
        monkeypatch.delenv(key, raising=False)


@pytest.fixture
def process_env(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'process.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    sessions = sessionmaker(engine, autoflush=False, expire_on_commit=False)

    def session_override():
        with sessions() as session:
            yield session

    app = FastAPI()
    app.include_router(data_router)
    app.include_router(process_router)
    app.include_router(ai_router)
    app.dependency_overrides[get_session] = session_override
    with TestClient(app) as client:
        yield client
    engine.dispose()


class StubLLM:
    def __init__(self, payloads):
        self.payloads = list(payloads)
        self.calls = 0
        self.messages = []
        self.model = "stub"

    async def complete(self, *, messages, schema, timeout):
        del schema, timeout
        self.calls += 1
        self.messages = messages
        payload = self.payloads.pop(0)
        if isinstance(payload, Exception):
            raise payload
        return payload


def _approval_payload():
    return GenDraft.model_validate({
        "isProcess": True,
        "suggestedName": "请假审批",
        "summary": "请假审批",
        "steps": [
            {"key": "s1", "kind": "start", "name": "开始", "next": "s2"},
            {"key": "s2", "kind": "task", "name": "提交申请", "owner": "张三", "next": "s3",
             "description": "核对订单中心",
             "assetRefs": [
                 {"ref": "a1", "mention": "订单中心", "confidence": 0.95},
                 {"ref": "a999", "mention": "虚构系统", "confidence": 0.99},
             ],
             "unmatchedMentions": ["外部短信网关", "幽灵系统"]},
            {"key": "s3", "kind": "decision", "name": "是否通过",
             "yes": {"to": "s4", "label": "通过"}, "no": {"to": "s5", "label": "驳回"}},
            {"key": "s4", "kind": "task", "name": "执行办理", "owner": "李四", "next": "s2"},
            {"key": "s5", "kind": "end", "name": "结束"},
        ],
    }).model_dump()


def _post(client, **fields):
    files = {}
    for key, value in fields.items():
        if key == "file":
            files["file"] = value
        elif value is not None:
            files[key] = (None, value)
    return client.post("/api/process/ai/draft", files=files)


def _snapshot(client):
    return client.get("/api/diagrams").json(), client.get("/api/export/json").json()


def _reachable(nodes, edges):
    start = next(node["id"] for node in nodes if node["type"] == "start")
    adjacent: dict[str, list[str]] = {}
    for edge in edges:
        adjacent.setdefault(edge["source"], []).append(edge["target"])
    seen: set[str] = set()
    stack = [start]
    while stack:
        current = stack.pop()
        if current in seen:
            continue
        seen.add(current)
        stack.extend(adjacent.get(current, []))
    return seen


def test_draft_validates_and_does_not_write(process_env, monkeypatch):
    client = process_env
    existing = _create_diagram(client)
    asset = client.post("/api/nodes", json={"id": "sys:订单中心", "name": "订单中心", "type": "system"})
    assert asset.status_code == 200, asset.text
    vanished = client.post("/api/nodes", json={"id": "sys:幽灵系统", "name": "幽灵系统", "type": "system"})
    assert vanished.status_code == 200, vanished.text
    hidden = client.patch("/api/nodes/sys:幽灵系统", json={"vanished": True})
    assert hidden.status_code == 200, hidden.text
    before = _snapshot(client)
    stub = StubLLM([_approval_payload()])
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: stub)

    response = _post(client, text=f"  {SOURCE}  ", name="请假审批", inputKind="description")

    assert response.status_code == 200, response.text
    body = response.json()
    draft = ProcessCreate.model_validate(body["draft"])
    assert body["sourceText"] == SOURCE
    assert body["stats"]["model"] == "stub"
    assert body["stats"]["llmCalls"] == 1
    assert body["stats"]["repaired"] is False
    assert isinstance(body["stats"]["latencyMs"], int)
    assert len([node for node in draft.nodes if node.type == "start"]) == 1
    assert any(node.type == "end" for node in draft.nodes)
    assert _reachable([node.model_dump() for node in draft.nodes], [edge.model_dump() for edge in draft.edges]) == {
        node.id for node in draft.nodes
    }
    assert all(node.assetIds == [] and node.x == 0 and node.y == 0 for node in draft.nodes)
    assert _snapshot(client) == before
    assert client.get(f"/api/diagrams/{existing['id']}").json() == existing
    prompt = json.dumps(stub.messages, ensure_ascii=False)
    assert "sys:订单中心" not in prompt
    assert "a1" in prompt
    available = {item["id"] for item in client.get("/api/assets").json()}
    assert "sys:幽灵系统" not in available
    assert {item["assetId"] for item in body["assetSuggestions"]} <= available
    assert any(item["assetId"] == "sys:订单中心" and item["preselect"] is True for item in body["assetSuggestions"])
    assert "a999" not in {item["assetId"] for item in body["assetSuggestions"]}
    assert "虚构系统" in {item["mention"] for item in body["unmatchedAssets"]}
    assert "幽灵系统" in {item["mention"] for item in body["unmatchedAssets"]}

    saved_body = body["draft"]
    names = [node["name"] for node in saved_body["nodes"]]
    created = client.post("/api/diagrams", json=saved_body)
    assert created.status_code == 200, created.text
    created_body = created.json()
    assert [node["name"] for node in created_body["nodes"]] == names
    assert all(node["assetIds"] == [] for node in created_body["nodes"])
    assert created_body["id"] != existing["id"]
    assert client.get(f"/api/diagrams/{existing['id']}").json() == existing
    listed = client.get("/api/diagrams").json()
    assert len(listed) == 2

    accepted = json.loads(json.dumps(saved_body))
    target = next(node for node in accepted["nodes"] if node["name"] == "提交申请")
    target["assetIds"] = ["sys:订单中心"]
    linked = client.post("/api/diagrams", json=accepted)
    assert linked.status_code == 200, linked.text
    usage = client.get("/api/assets/sys:%E8%AE%A2%E5%8D%95%E4%B8%AD%E5%BF%83/diagrams").json()
    assert any(item["id"] == linked.json()["id"] for item in usage)
    stale = client.put(f"/api/diagrams/{existing['id']}", json={
        "name": existing["name"], "description": existing["description"],
        "nodes": existing["nodes"], "edges": existing["edges"], "expected_revision": "stale",
    })
    assert stale.status_code == 409
    assert client.get(f"/api/diagrams/{existing['id']}").json() == existing


def test_not_a_process_writes_nothing(process_env, monkeypatch):
    client = process_env
    before = _snapshot(client)
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: StubLLM([{
        "isProcess": False, "notProcessReason": "这是一段产品介绍", "steps": [],
    }]))
    response = _post(client, text="这是一段产品介绍，没有办理步骤。")
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "not_a_process"
    assert detail["message"].startswith(NOT_A_PROCESS)
    assert PARSE_FAILED not in detail["message"]
    assert "扫描" not in detail["message"]
    assert _snapshot(client) == before


def test_repair_then_success_still_writes_nothing(process_env, monkeypatch):
    client = process_env
    before = _snapshot(client)
    stub = StubLLM([{"not": "a draft"}, _approval_payload()])
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: stub)
    response = _post(client, text=SOURCE)
    assert response.status_code == 200, response.text
    assert response.json()["stats"]["llmCalls"] == 2
    assert response.json()["stats"]["repaired"] is True
    assert stub.calls == 2
    assert _snapshot(client) == before


def test_unrepairable_model_output_is_502(process_env, monkeypatch):
    client = process_env
    before = _snapshot(client)
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: StubLLM([{"not": "a draft"}, {"still": "bad"}]))
    response = _post(client, text=SOURCE)
    assert response.status_code == 502
    assert response.json()["detail"]["code"] == "llm_failed"
    assert _snapshot(client) == before


def test_timeout_is_504_without_waiting_for_the_full_budget(process_env, monkeypatch):
    client = process_env

    class Slow:
        model = "slow"

        async def complete(self, *, messages, schema, timeout):
            del messages, schema, timeout
            await asyncio.sleep(5)

    monkeypatch.setattr("app.process_ai.api.BUDGET_SECONDS", 0.2)
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: Slow())
    started = time.monotonic()
    response = _post(client, text=SOURCE)
    assert response.status_code == 504
    assert response.json()["detail"]["code"] == "llm_timeout"
    assert time.monotonic() - started < 2


def test_missing_llm_config_is_502_and_does_not_call_http(process_env, monkeypatch):
    client = process_env

    class Boom:
        def __init__(self, *args, **kwargs):
            raise AssertionError("不应访问网络")

    monkeypatch.setattr("app.process_ai.llm.httpx.AsyncClient", Boom)
    response = _post(client, text=SOURCE)
    assert response.status_code == 502
    detail = response.json()["detail"]
    assert detail["code"] == "llm_failed"
    assert "未配置语言模型" in detail["message"]
    assert "key" not in detail["message"].lower()


def test_fixture_env_generates_without_a_key(process_env, monkeypatch, tmp_path):
    client = process_env
    path = tmp_path / "fixture.json"
    path.write_text(json.dumps(_approval_payload()), encoding="utf-8")
    monkeypatch.setenv("XILIU_LLM_FIXTURE", str(path))
    before = _snapshot(client)
    response = _post(client, text=SOURCE, matchAssets="false")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["stats"]["model"] == "fixture"
    assert body["assetSuggestions"] == []
    assert all(node["assetIds"] == [] for node in body["draft"]["nodes"])
    assert _snapshot(client) == before


def test_file_errors_and_docx_pdf_round_trip(process_env, monkeypatch):
    client = process_env
    stub = StubLLM([_approval_payload(), _approval_payload(), _approval_payload()])
    monkeypatch.setattr("app.process_ai.api.build_client", lambda: stub)
    before = _snapshot(client)

    scanned = _post(client, file=("scan.pdf", _blank_pdf(), "application/pdf"))
    assert scanned.status_code == 422
    assert scanned.json()["detail"]["code"] == "scanned_pdf"
    assert scanned.json()["detail"]["message"] == SCANNED_PDF

    corrupt = _post(client, file=("bad.pdf", b"not a pdf", "application/pdf"))
    assert corrupt.status_code == 422
    assert corrupt.json()["detail"]["code"] == "parse_failed"
    assert corrupt.json()["detail"]["message"] == PARSE_FAILED
    assert corrupt.json()["detail"]["message"] != scanned.json()["detail"]["message"]

    unsupported = _post(client, file=("pic.png", b"abc", "image/png"))
    assert unsupported.status_code == 415
    assert unsupported.json()["detail"]["code"] == "unsupported_type"

    empty = _post(client, file=("empty.txt", b" \n\t", "text/plain"))
    assert empty.status_code == 422 and empty.json()["detail"]["code"] == "empty_text"

    too_long = _post(client, text="字" * 20001)
    assert too_long.status_code == 422 and too_long.json()["detail"]["code"] == "text_too_long"
    assert stub.calls == 0

    huge = _post(client, file=("big.txt", b"a" * (10 * 1024 * 1024 + 1), "text/plain"))
    assert huge.status_code == 413 and huge.json()["detail"]["code"] == "file_too_large"

    gbk = "会议决定由王敏负责审批。张三提交申请后，判断是否通过。通过则由李四执行办理，驳回则结束。"
    decoded = _post(client, file=("minutes.txt", gbk.encode("gbk"), "text/plain"), inputKind="minutes")
    assert decoded.status_code == 200, decoded.text
    assert "王敏" in decoded.json()["sourceText"]
    assert "王敏" in json.dumps(stub.messages, ensure_ascii=False)

    from docx import Document
    document = Document()
    document.add_paragraph("第一步 提交申请")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "负责人"
    table.rows[0].cells[1].text = "张三"
    buffer = io.BytesIO()
    document.save(buffer)
    docx = _post(client, file=("sop.docx", buffer.getvalue(),
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                 inputKind="sop")
    assert docx.status_code == 200, docx.text
    assert "第一步 提交申请" in docx.json()["sourceText"]
    assert "负责人 | 张三" in docx.json()["sourceText"]

    pdf = _post(client, file=("leave.pdf", _pdf_with_text("Applicant submits the leave request for approval today"),
                              "application/pdf"))
    assert pdf.status_code == 200, pdf.text
    assert "leave request" in pdf.json()["sourceText"]
    assert _snapshot(client) == before


def test_input_xor_and_invalid_options(process_env):
    client = process_env
    both = _post(client, text="有文字", file=("a.txt", b"file", "text/plain"))
    assert both.status_code == 422 and both.json()["detail"]["code"] == "empty_text"
    neither = _post(client, inputKind="description")
    assert neither.status_code == 422 and neither.json()["detail"]["code"] == "empty_text"
    bad_kind = _post(client, text="有文字", inputKind="novel")
    assert bad_kind.status_code == 422
    bad_match = _post(client, text="有文字", matchAssets="maybe")
    assert bad_match.status_code == 422


def _create_diagram(client):
    payload = {
        "name": "已有流程", "description": "保持不变",
        "nodes": [{
            "id": "step", "type": "task", "name": "步骤", "x": 1.0, "y": 2.0,
            "owner": "", "description": "", "input": "", "output": "", "assetIds": [],
        }],
        "edges": [],
    }
    response = client.post("/api/diagrams", json=payload)
    assert response.status_code == 200, response.text
    return response.json()
