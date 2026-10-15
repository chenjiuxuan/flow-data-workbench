"""文件与文字解析，对应验收 12。"""
import asyncio
import io

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.process_ai.errors import PARSE_FAILED, SCANNED_PDF, DraftError
from app.process_ai.extract import (
    MAX_FILE_BYTES,
    extract_bytes,
    normalize_text,
    prepare_text,
    read_limited,
)


def test_normalize_and_length_limits():
    assert normalize_text("a \t b\r\n\r\n\r\nc") == "a b\n\nc"
    assert prepare_text("字" * 20000) == "字" * 20000
    with pytest.raises(DraftError) as too_long:
        prepare_text("字" * 20001)
    assert too_long.value.status == 422 and too_long.value.code == "text_too_long"
    with pytest.raises(DraftError) as empty:
        prepare_text(" \n\t")
    assert empty.value.code == "empty_text"


def test_txt_and_md_utf8_bom_and_gbk():
    text = "会议决定由王敏负责审批"
    assert extract_bytes("notes.txt", text.encode("utf-8")) == text
    assert extract_bytes("NOTES.MD", text.encode("utf-8-sig")) == text
    assert extract_bytes("minutes.txt", text.encode("gbk")) == text
    assert extract_bytes("empty.txt", b"") == ""


def test_docx_keeps_paragraph_and_table_order():
    document = Document()
    document.add_paragraph("第一步 提交")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "负责人"
    table.rows[0].cells[1].text = "王敏"
    table.rows[1].cells[0].text = "系统"
    table.rows[1].cells[1].text = "订单中心"
    document.add_paragraph("第二步 归档")
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_bytes("sop.DOCX", buffer.getvalue())
    assert text.index("第一步 提交") < text.index("负责人 | 王敏") < text.index("系统 | 订单中心")
    assert text.index("系统 | 订单中心") < text.index("第二步 归档")


def test_pdf_text_layer_and_scanned_pdf_are_distinct():
    text = extract_bytes("leave.pdf", _pdf_with_text("Applicant submits the leave request for approval today"))
    assert "leave request" in text
    with pytest.raises(DraftError) as scanned:
        extract_bytes("scan.pdf", _blank_pdf())
    assert scanned.value.status == 422
    assert scanned.value.code == "scanned_pdf"
    assert scanned.value.message == SCANNED_PDF
    assert scanned.value.message == "这个 PDF 像是扫描件，读不到文字。请把里面的文字复制进来，或上传 Word/文本版"
    with pytest.raises(DraftError) as broken:
        extract_bytes("broken.pdf", b"not a pdf")
    assert broken.value.code == "parse_failed"
    assert broken.value.message == PARSE_FAILED
    assert broken.value.message != scanned.value.message


def test_corrupt_docx_and_unsupported_type():
    with pytest.raises(DraftError) as corrupt:
        extract_bytes("bad.docx", b"this is not a docx")
    assert corrupt.value.code == "parse_failed"
    with pytest.raises(DraftError) as unsupported:
        extract_bytes("picture.png", b"abc")
    assert unsupported.value.status == 415 and unsupported.value.code == "unsupported_type"
    with pytest.raises(DraftError) as old_word:
        extract_bytes("legacy.doc", b"abc")
    assert old_word.value.code == "unsupported_type"


def test_read_limited_stops_past_ten_mib():
    async def scenario():
        upload = _Chunks([b"x" * (1024 * 1024)] * 12)
        with pytest.raises(DraftError) as too_big:
            await read_limited(upload)
        assert too_big.value.status == 413 and too_big.value.code == "file_too_large"
        assert upload.reads == 11

        exact = _Chunks([b"y" * (1024 * 1024)] * 10 + [b""])
        data = await read_limited(exact)
        assert len(data) == MAX_FILE_BYTES

        declared = _Chunks([b"z" * 10])
        declared.size = MAX_FILE_BYTES + 1
        with pytest.raises(DraftError) as declared_error:
            await read_limited(declared)
        assert declared_error.value.code == "file_too_large"
        assert declared.reads == 0

    asyncio.run(scenario())


class _Chunks:
    def __init__(self, chunks):
        self._chunks = list(chunks)
        self.reads = 0
        self.size = None
        self.filename = "big.txt"

    async def read(self, size=-1):
        del size
        self.reads += 1
        if not self._chunks:
            return b""
        return self._chunks.pop(0)


def _blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pdf_with_text(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 120 Td ({text}) Tj ET".encode("ascii"))
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Contents")] = writer._add_object(stream)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
