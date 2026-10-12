"""Excel 台账导入：提供模板下载与解析。

模板列（ sheet「数据流台账」）：
源系统 | 源系统类型 | 目标系统 | 目标系统类型 | 流转方式 | 数据内容 | 字段映射 | 频率 | 负责人 | 备注

- 源/目标系统 → 系统节点（外部方类型可选）；
- 数据内容形如 库.表 时额外建表节点，并用 产出/消费 两条边挂到系统上；
- 每行生成一条 系统→系统 数据流。
"""
import io
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError as XMLParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .ids import flow_id, sys_id, table_id
from .merge import EdgeDraft, NodeDraft

TEMPLATE_SHEET = "数据流台账"
HEADERS = ["源系统", "源系统类型(系统/外部方)", "目标系统", "目标系统类型(系统/外部方)",
           "流转方式(接口/库表同步/文件/人工)", "数据内容(库.表 或 说明)",
           "字段映射(源字段->目标字段,逗号分隔)", "频率(实时/小时/日/月)", "负责人", "备注"]

_TABLE_PATTERN = re.compile(r"^[A-Za-z\u4e00-\u9fa5_][\w\u4e00-\u9fa5]*\.[A-Za-z_][\w]*$")
MAX_EXCEL_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_EXCEL_ZIP_ENTRIES = 2048
MAX_EXCEL_ROWS = 100000


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_SHEET
    ws.append(HEADERS)
    ws.append(["订单中心", "系统", "支付中心", "系统", "接口", "订单支付结果", "order_no->order_no",
               "实时", "张三", "示例行，可删除"])
    ws.append(["订单库", "系统", "数据仓库", "系统", "库表同步", "orders.orders", "order_id->order_id,amount->amount",
               "日", "李四", "每日凌晨 T+1"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = [16, 18, 16, 18, 20, 26, 34, 16, 10, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(value) -> str:
    return str(value).strip() if value is not None else ""


def parse_excel(content: bytes) -> tuple[list[NodeDraft], list[EdgeDraft], list[str]]:
    """返回 (节点草稿, 边草稿, 错误列表)。"""
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > MAX_EXCEL_ZIP_ENTRIES or
                    sum(entry.file_size for entry in entries) > MAX_EXCEL_EXPANDED_BYTES):
                raise ValueError("Excel 解压后过大或包含过多文件，最大支持 64 MiB / 2048 个条目")
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, XMLParseError, RuntimeError) as exc:
        raise ValueError("无效或损坏的 Excel 文件，请使用下载的模板") from exc
    try:
        if TEMPLATE_SHEET not in wb.sheetnames:
            raise ValueError(f"缺少工作表「{TEMPLATE_SHEET}」，请使用下载的模板填写")
        ws = wb[TEMPLATE_SHEET]
        # 部分生成工具会低报/夸大 dimension；否则 read_only 可能漏掉所有真实行，
        # 把正常台账误判成空快照。按 XML 实际内容读取并应用自己的行数限制。
        ws.reset_dimensions()
        rows = []
        # 只读取模板所需列，且不信任工作表中可能被夸大的行/列维度。
        for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
            if len(rows) >= MAX_EXCEL_ROWS:
                raise ValueError("Excel 超过最大支持的 100000 行，请拆分导入")
            rows.append(row)
    except (BadZipFile, OSError, XMLParseError, RuntimeError) as exc:
        raise ValueError("无效或损坏的 Excel 文件，请使用下载的模板") from exc
    finally:
        wb.close()

    nodes: dict[str, NodeDraft] = {}
    edges: dict[str, EdgeDraft] = {}
    errors: list[str] = []

    for index, row in enumerate(rows, start=2):
        values = [_cell_str(v) for v in (list(row) + [""] * len(HEADERS))[:len(HEADERS)]]
        src_name, src_kind, dst_name, dst_kind, method, content, fields, freq, owner, note = values
        if not src_name and not dst_name:
            continue
        if not src_name or not dst_name:
            errors.append(f"第 {index} 行：源系统与目标系统都要填写，已跳过")
            continue

        src_id = sys_id(src_name)
        dst_id = sys_id(dst_name)
        nodes[src_id] = NodeDraft(src_id, "external" if "外部" in src_kind else "system", src_name)
        nodes[dst_id] = NodeDraft(dst_id, "external" if "外部" in dst_kind else "system", dst_name)

        attrs = {"method": method, "frequency": freq, "owner": owner, "note": note, "content": content}
        # 空单元格是明确清空；省略键会让 merge 错误地保留上次的自动映射。
        attrs["fields"] = [f.strip() for f in re.split(r"[,;，；]", fields) if f.strip()]

        edge = EdgeDraft(
            id=flow_id(src_id, dst_id, content or method or "data"),
            source_id=src_id, target_id=dst_id, type="flow",
            label=method or "数据流", attrs=attrs,
        )
        if edge.id in edges:
            edges[edge.id].attrs = attrs  # 同一 流向+内容 视为同一流转
        else:
            edges[edge.id] = edge

        # 数据内容是 库.表 时，建表节点挂在两个系统之间
        if content and _TABLE_PATTERN.match(content):
            tid = table_id(*content.split(".", 1))
            display = content.lower()
            nodes[tid] = NodeDraft(tid, "table", display,
                                   attrs={"database": content.split(".", 1)[0].lower(),
                                          "table": content.split(".", 1)[1].lower()})
            produce = EdgeDraft(id=f"produces:{src_id}->{tid}", source_id=src_id, target_id=tid,
                                type="flow", label="产出", attrs={"relation": "produces"})
            consume = EdgeDraft(id=f"consumes:{tid}->{dst_id}", source_id=tid, target_id=dst_id,
                                type="flow", label="消费", attrs={"relation": "consumes"})
            edges.setdefault(produce.id, produce)
            edges.setdefault(consume.id, consume)

    return list(nodes.values()), list(edges.values()), errors
